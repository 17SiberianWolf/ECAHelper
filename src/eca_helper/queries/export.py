"""Excel / CSV 中英双语导出（系统设计书.md §5.3 / T12 / Q7）。

导出范围 = 当前查询筛选（复用 aggregate 的 WHERE 构造 + 排除条款）。
表头使用 config.bilingual_header 生成中英双语；task_zh 为可选业务枚举中文映射列。

文件名约定（决策 +3）：导出_<类型>_<YYYYMMDD_HHMM>.<xlsx|csv>
"""

from __future__ import annotations

import csv
import datetime
import re
from pathlib import Path

from config import (
    BUSINESS_ENUM_ZH,
    EXPORT_COLUMNS,
    EXPORT_DIR,
    EXPORT_PREFIX,
    bilingual_header,
)
from eca_helper.db import get_connection
from eca_helper.queries.aggregate import (
    EXCLUSION_SQL,
    PROJECT_BASE_SQL,
    QueryFilter,
    _base_where,
    _empty_project_cond,
)

# 非 DB 计算列（task_zh）
_DERIVED = {"task_zh"}


def _select_columns() -> list[str]:
    return [c for c in EXPORT_COLUMNS if c not in _DERIVED]


def fetch_rows(conn, f: QueryFilter) -> list[dict]:
    """按筛选条件取明细行（含排除条款）。

    导出列**按列定制**（决策 9 / §5.4）：
      - `project_id_norm` -> PROJECT_BASE_SQL AS project_id_norm（**父号主列**，归并口径）
      - `project_id_raw`  -> r.project_id_raw AS project_id_raw（**原始项目号**，既有列）
      - 其余 -> f"r.{c}"
    父号列与原始号列**同一行并存**（不新增 DB 列）。
    """
    params: dict = {}
    where = _base_where(f, params) + _empty_project_cond(f)
    exprs: list[str] = []
    for c in _select_columns():  # _select_columns() 已剔除派生列 task_zh
        if c == "project_id_norm":
            exprs.append(PROJECT_BASE_SQL + " AS project_id_norm")
        elif c == "project_id_raw":
            exprs.append("r.project_id_raw AS project_id_raw")
        else:
            exprs.append(f"r.{c}")
    sql = (
        "SELECT " + ", ".join(exprs)
        + " FROM timesheet_record r WHERE 1=1 " + EXCLUSION_SQL + where
        + " ORDER BY r.report_month, r.source_file, r.source_row"
    )
    cur = conn.cursor()
    cur.execute(sql, params)
    rows = [dict(x) for x in cur.fetchall()]
    for r in rows:
        r["task_zh"] = BUSINESS_ENUM_ZH.get(r.get("task") or "", "")
    return rows


def _headers(lang: str) -> list[str]:
    headers: list[str] = []
    for col in EXPORT_COLUMNS:
        if col == "task_zh":
            headers.append("Task (中文)" if lang != "zh" else "任务(中文)")
        else:
            headers.append(bilingual_header(col, lang))
    return headers


def _row_values(row: dict) -> list:
    return [row.get(c, "") for c in EXPORT_COLUMNS]


def build_csv(rows: list[dict], lang: str) -> str:
    import io

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_headers(lang))
    for r in rows:
        writer.writerow(_row_values(r))
    return buf.getvalue()


def build_workbook(rows: list[dict], lang: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    headers = _headers(lang)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(_row_values(r))
    # 粗略列宽
    from openpyxl.utils import get_column_letter

    for c in range(1, len(headers) + 1):
        col_name = EXPORT_COLUMNS[c - 1]
        sample = [len(str(r.get(col_name, ""))) for r in rows[:500]]
        max_len = max([len(str(headers[c - 1]))] + sample) if sample else len(str(headers[c - 1]))
        ws.column_dimensions[get_column_letter(c)].width = min(max(max_len + 2, 10), 40)
    ws.freeze_panes = "A2"
    return wb


def _safe_title(title: str) -> str:
    t = (title or "数据").strip()
    t = re.sub(r"[\\/:*?\"<>|]", "_", t)
    return t[:30] or "数据"


def export_to_file(conn, f: QueryFilter, fmt: str = "xlsx", lang: str = "both",
                   title: str = "数据") -> tuple[str, str]:
    """导出到 EXPORT_DIR，返回 (绝对路径, 文件名)。"""
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    rows = fetch_rows(conn, f)
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    fname = f"{EXPORT_PREFIX}_{_safe_title(title)}_{ts}.{fmt}"
    path = EXPORT_DIR / fname

    if fmt == "csv":
        # csv.writer 默认行尾为 "\r\n"；若用 write_text（newline 默认转换）
        # 会在 Windows 上二次转换成 "\r\r\n"，导致 Excel 打开后数据行之间夹空行。
        # 故直接写字节，绕过换行转换，保留 csv.writer 原始的 "\r\n"。
        content = build_csv(rows, lang)
        path.write_bytes(content.encode("utf-8-sig"))
    else:
        wb = build_workbook(rows, lang)
        wb.save(str(path))
    return str(path), fname


def export_from_request(conn, payload: dict):
    """从 API 请求体构造 QueryFilter 并导出。"""
    f = QueryFilter(
        report_month_start=payload.get("start"),
        report_month_end=payload.get("end"),
        organization=payload.get("organization"),
        cost_center=payload.get("cost_center"),
        task=payload.get("task"),
        resource_id=payload.get("resource_id"),
        project_id=payload.get("project_id"),
        include_empty_project=bool(payload.get("include_empty_project", False)),
    )
    fmt = payload.get("format", "xlsx")
    lang = payload.get("lang", "both")
    title = payload.get("title", "数据")
    if fmt not in ("xlsx", "csv"):
        fmt = "xlsx"
    if lang not in ("both", "en", "zh"):
        lang = "both"
    return export_to_file(conn, f, fmt=fmt, lang=lang, title=title)
