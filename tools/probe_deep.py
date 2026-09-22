import collections
import glob
import io
import os
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]  # tools/ 的上一级 = 项目根
SRC = ROOT / "OriginSource"
out = io.StringIO()
def w(*a): print(*a, file=out)
def norm(s):
    if s is None: return ""
    return re.sub(r"\s+", " ", str(s).replace("\xa0"," ")).strip()

# ---- A. 在管项目 file ----
f = os.path.join(SRC, "在管项目2026.08.xlsx")
wb = load_workbook(f, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
w("=== 在管项目2026.08.xlsx  dims:", ws.calculate_dimension(), "===")
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i > 20: break
    w(f"  r{i}: {[norm(c) for c in r][:12]}")
wb.close()

# ---- B. resource id -> name variants ----
files = sorted([x for x in glob.glob(os.path.join(SRC,"*.xlsx")) if not os.path.basename(x).startswith("~$")])
id2names = collections.defaultdict(set)
name2ids = collections.defaultdict(set)
era_name_sample = {"v1": set(), "v2": set()}
dupkeys = {}
zero_neg = 0
total_h = 0.0
strhours = []
blank_projectid = 0
rows_total = 0

for fp in files:
    base = os.path.basename(fp)
    wb = load_workbook(fp, read_only=True, data_only=True)
    target = None
    for sn in wb.sheetnames:
        ws = wb[sn]
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i > 5: break
            cells = [norm(c) for c in r]
            if any(c=="Resourcename" for c in cells) and any("Actuals Total in h" in c for c in cells):
                target = (sn, i, cells); break
        if target: break
    if not target:
        w("skip (no data sheet):", base); wb.close(); continue
    sn, hi, hdr = target
    while hdr and hdr[-1]=="": hdr.pop()
    era = "v2" if "WH Cost Center" in hdr else "v1"
    ri = hdr.index("Resourcename"); ii = hdr.index("Resource ID")
    pi = hdr.index("Project ID"); vi = hdr.index("Actuals Total in h")
    ws = wb[sn]
    seen = collections.Counter()
    for r in ws.iter_rows(min_row=hi+2, values_only=True):
        cells = [norm(c) for c in r]
        nm, rid = cells[ri], cells[ii]
        pid = cells[pi]
        v = r[vi] if vi < len(r) else None
        if nm=="" and rid=="" and pid=="" and (v is None or v==""): continue
        rows_total += 1
        if nm: id2names[rid].add(nm); name2ids[nm].add(rid)
        if len(era_name_sample[era]) < 4000 and nm: era_name_sample[era].add(nm)
        if pid == "": blank_projectid += 1
        key = tuple(cells[:8]) + (str(v),)
        seen[key] += 1
        if isinstance(v,(int,float)):
            total_h += v
            if v <= 0: zero_neg += 1
        elif v not in (None,""):
            strhours.append((base, str(v)))
    d = sum(c-1 for c in seen.values() if c>1)
    if d: dupkeys[base] = d
    wb.close()

w("\n=== 资源ID -> 姓名 变体（一个人多个名字写法）===")
multi = {k:v for k,v in id2names.items() if len(v)>1}
w("有多个姓名写法的 Resource ID 数量:", len(multi), "/ 总ID数:", len(id2names))
for k,v in list(multi.items())[:15]: w(f"   {k}: {sorted(v)}")

w("\n=== 姓名 -> 多个 Resource ID ===")
mn = {k:v for k,v in name2ids.items() if len(v)>1}
w("有多个ID的姓名数量:", len(mn))
for k,v in list(mn.items())[:10]: w(f"   {k}: {sorted(v)}")

w("\n=== 校验 ===")
w("总数据行:", rows_total)
w("累积工时 sum(Actuals Total in h) =", round(total_h,2))
w("<=0 的工时行数:", zero_neg)
w("Project ID 为空的行数:", blank_projectid)
w("字符串型工时样例:", strhours[:10], "共", len(strhours))
w("同文件内重复行(疑似重复数据) 文件数:", len(dupkeys))
for k,v in list(dupkeys.items())[:10]: w("   ", k, "重复行数:", v)

w("\n=== 人名写法对比 ===")
w("v1(旧) 样例:", sorted(era_name_sample["v1"])[:12])
w("v2(新) 样例:", sorted(era_name_sample["v2"])[:12])
w("v1 人数:", len(era_name_sample["v1"]), " v2 人数:", len(era_name_sample["v2"]))
w("两版共有姓名:", len(era_name_sample["v1"] & era_name_sample["v2"]))

with open(ROOT / "tools" / "_deep_report.txt", "w", encoding="utf-8") as fh:
    fh.write(out.getvalue())
print("ok")
