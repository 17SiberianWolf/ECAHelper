import collections
import datetime
import glob
import io
import os
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]  # tools/ 的上一级 = 项目根
SRC = ROOT / "OriginSource"
out = io.StringIO()
def w(*a): print(*a, file=out)

files = sorted([f for f in glob.glob(os.path.join(SRC, "*.xlsx")) if not os.path.basename(f).startswith("~$")])

def norm(s):
    if s is None: return ""
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip()

rows_report = []
all_headers = collections.Counter()
org_counter = collections.Counter()
valtype_counter = collections.Counter()
datecol_notes = []

for f in files:
    base = os.path.basename(f)
    wb = load_workbook(f, read_only=True, data_only=True)
    found = None
    for sn in wb.sheetnames:
        ws = wb[sn]
        maxr = min(ws.max_row or 0, 12)
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i > maxr: break
            cells = [norm(c) for c in r]
            if any(c == "Resourcename" for c in cells) and any("Actuals Total in h" in c for c in cells):
                found = (sn, i, cells, ws)
                break
        if found: break
    if not found:
        w(f"!! NO DATA SHEET: {base} | sheets={wb.sheetnames}")
        wb.close()
        continue
    sn, hdr_i, hdr, ws = found
    # clean trailing empties
    while hdr and hdr[-1] == "": hdr.pop()
    all_headers[tuple(hdr)] += 1

    # count data rows + value types + distinct orgs
    nrows = 0
    tcount = collections.Counter()
    totrows = []
    for j, r in enumerate(ws.iter_rows(min_row=hdr_i + 2, values_only=True)):
        cells = [norm(c) for c in r]
        if all(c == "" for c in cells): continue
        first = cells[0] if cells else ""
        if first.lower().startswith(("grand total", "total", "sum of", "row labels")):
            totrows.append(first)
            continue
        nrows += 1
        idx = hdr.index("Actuals Total in h") if "Actuals Total in h" in hdr else 8
        v = r[idx] if idx < len(r) else None
        if isinstance(v, (int, float)): tcount["num"] += 1
        elif v is None: tcount["none"] += 1
        else: tcount["str:" + str(v)[-1]] += 1
        if "Organization" in hdr:
            oi = hdr.index("Organization")
            if oi < len(cells): org_counter[cells[oi]] += 1
    w(f"{base}\n   sheet='{sn}' headerRow={hdr_i+1} cols={len(hdr)} dataRows={nrows} totalRows={totrows} vals={dict(tcount)}")
    w(f"   header={hdr}")
    wb.close()

w("\n\n=== DISTINCT HEADER SCHEMAS ===")
for h, c in all_headers.most_common():
    w(f"[{c} files] {list(h)}")

w("\n\n=== ORGANIZATION values ===")
for o, c in org_counter.most_common():
    w(f"  '{o}': {c}")

with open(ROOT / "tools" / "_scan_report.txt", "w", encoding="utf-8") as fh:
    fh.write(out.getvalue())
print("ok")
